import openpyxl

wb = openpyxl.load_workbook(r'd:\PowerBI\archive\IPL 2025 Auction Data.xlsx')

print("Sheet1 (Auction Data):")
s1 = wb['Sheet1']
for r in list(s1.iter_rows(values_only=True))[:10]:
    print(r)

print("\nSheet2 (Performance Data):")
s2 = wb['Sheet2']
for r in list(s2.iter_rows(values_only=True))[:10]:
    print(r)
