import csv
import openpyxl

# Load workbook
wb_path = r'd:\PowerBI\archive\IPL 2025 Auction Data.xlsx'
wb = openpyxl.load_workbook(wb_path)
sheet2 = wb['Sheet2']

# Load data from CSV files
name_map = {
    'Suryakumar Yadav': 'Surya Kumar Yadav',
    'KL Rahul': 'K L Rahul',
    'Varun Chakravarthy': 'Varun Chakaravarthy'
}

batters = {}
with open(r'd:\PowerBI\archive\IPL2025Batters.csv', encoding='utf-8-sig') as f:
    r = csv.reader(f)
    next(r)
    for row in r:
        if row:
            batters[row[0].strip()] = row

bowlers = {}
with open(r'd:\PowerBI\archive\IPL2025Bowlers.csv', encoding='utf-8-sig') as f:
    r = csv.reader(f)
    next(r)
    for row in r:
        if row:
            bowlers[row[0].strip()] = row

# We will update the rows in Sheet2
# Header is at row 1: Player, Team, Role, Matches, Runs, StrikeRate, Average, Wickets, Economy
# Let's map column indexes (1-based for openpyxl)
col_map = {
    'Player': 1,
    'Team': 2,
    'Role': 3,
    'Matches': 4,
    'Runs': 5,
    'StrikeRate': 6,
    'Average': 7,
    'Wickets': 8,
    'Economy': 9
}

# CSV output data
csv_rows = [["Player", "Team", "Role", "Matches", "Runs", "StrikeRate", "Average", "Wickets", "Economy", "Source"]]

# Iterate through sheet2 rows starting from row 2
for row_idx in range(2, sheet2.max_row + 1):
    player_name = sheet2.cell(row=row_idx, column=col_map['Player']).value
    if not player_name:
        continue
    
    player_name = player_name.strip()
    mapped = name_map.get(player_name, player_name)
    
    b_row = batters.get(mapped)
    bo_row = bowlers.get(mapped)
    
    # Matches
    matches = 0
    if b_row:
        matches = int(b_row[3])
    elif bo_row:
        matches = int(bo_row[3])
        
    # Batting
    if b_row:
        runs = int(b_row[2])
        sr = float(b_row[9])
        avg = b_row[7]
        if avg != '-':
            avg = float(avg)
    else:
        runs = 0
        sr = 0.0
        avg = 0.0
        
    # Bowling
    if bo_row:
        wickets = int(bo_row[2])
        eco = float(bo_row[9])
    else:
        wickets = 0
        eco = 0.0
        
    # Update excel cells
    sheet2.cell(row=row_idx, column=col_map['Matches'], value=matches)
    sheet2.cell(row=row_idx, column=col_map['Runs'], value=runs)
    sheet2.cell(row=row_idx, column=col_map['StrikeRate'], value=sr)
    sheet2.cell(row=row_idx, column=col_map['Average'], value=avg)
    sheet2.cell(row=row_idx, column=col_map['Wickets'], value=wickets)
    sheet2.cell(row=row_idx, column=col_map['Economy'], value=eco)
    
    # Add to CSV row list
    source = f"IPL Official Website (Name Mismatch: '{mapped}')" if player_name in name_map else "IPL Official Website"
    csv_rows.append([
        player_name,
        sheet2.cell(row=row_idx, column=col_map['Team']).value,
        sheet2.cell(row=row_idx, column=col_map['Role']).value,
        str(matches),
        str(runs),
        str(sr),
        str(avg),
        str(wickets),
        str(eco),
        source
    ])

# Save Excel file
wb.save(wb_path)
print("Excel sheet updated successfully.")

# Save CSV file
csv_path = r'd:\PowerBI\archive\IPL_2025_Verified_Stats.csv'
with open(csv_path, 'w', newline='', encoding='utf-8') as f:
    writer = csv.writer(f)
    writer.writerows(csv_rows)
print("CSV file created successfully.")
